"""Write a tailored CV for one graded posting, and read the ones already written.

Generation is one Sonnet call for one posting -- slower and dearer than the
Haiku grading on page 4, which is why it is a per-row button here rather than a
batch run.
"""

import io
import json

import common
import pandas as pd
import streamlit as st
from docx import Document

from ai import cv_generation
from cv_generator import docx_gen, pdf_gen
from db import db_connection

ELIGIBLE_COLUMNS = ["job_id", "grade", "company", "title", "place", "url",
                    "analysed_at"]
CV_COLUMNS = ["cv_id", "title", "company", "grade", "locale", "cv", "url", "timestamp"]

#: Label -> SQLite modifier for datetime('now', ?). Same shape as the
#: AI analysis page's backlog window; "all time" is an epoch, not a special
#: case, so the query stays one shape.
WINDOWS = {
    "last 24 hours": "-24 hours",
    "last 7 days": "-7 days",
    "last 30 days": "-30 days",
    "all time": "-100 years",
}

GRADE_COLUMN = st.column_config.ProgressColumn("grade", min_value=0,
                                               max_value=100, format="%d")

st.set_page_config(page_title="CV generation", layout="wide")
st.title("CV generation")

db_connection.create_tables()

left, right = st.columns(2)
min_grade = left.slider("Minimum grade", 0, 100, 70,
                        help="Postings graded below this are not worth a "
                             "tailored CV, so they stay out of the list.")
window_label = right.selectbox("Graded within", list(WINDOWS), index=3,
                               help="Only show postings analysed within "
                                    "this window.")
window = WINDOWS[window_label]

eligible = pd.DataFrame(db_connection.select_jobs_for_cv(min_grade, window),
                        columns=ELIGIBLE_COLUMNS)
generated = pd.DataFrame(db_connection.select_generated_cvs(),
                         columns=CV_COLUMNS)

columns = st.columns(2)
columns[0].metric(f"Eligible postings (grade >= {min_grade})",
                  f"{len(eligible):,}")
columns[1].metric("CVs generated", f"{len(generated):,}")

st.subheader("Postings without a CV")
st.caption("Best fit first. Select one, then generate.")

table = st.dataframe(
    eligible[["grade", "company", "title", "place", "analysed_at", "url"]],
    use_container_width=True, hide_index=True, on_select="rerun",
    selection_mode="single-row",
    column_config={"url": st.column_config.LinkColumn("url"),
                   "grade": GRADE_COLUMN,
                   "analysed_at": st.column_config.DatetimeColumn(
                       "analysed", format="DD/MM/YYYY HH:mm")},
)

selected = table.selection["rows"]
job = eligible.iloc[selected[0]] if selected else None

st.caption("One Sonnet call: budget around half a minute.")
launched = st.button(
    f"Generate a CV for {job['title']}" if job is not None
    else "Generate a CV (select a posting first)",
    type="primary", disabled=job is None,
)

if launched:
    def workflow(log):
        """Generate and store the CV for the selected posting."""
        log(f"job {job['job_id']}: {job['company']} - {job['title']}")
        locale = cv_generation.generate_cv(int(job["job_id"]))
        log(f"stored, written in {locale}")

        return {"locale": locale}

    stats = common.run_with_log("generating the CV", workflow)

    if stats:
        st.success(f"CV stored, written in {stats['locale']}.")
        st.button("Reload the lists")

st.divider()
st.subheader("Generated CVs")

if generated.empty:
    st.info("No CVs generated yet.")
    st.stop()

st.caption("Newest first. Select one to read it.")

cv_table = st.dataframe(
    generated[["title", "company", "grade", "timestamp"]], use_container_width=True, hide_index=True,
    on_select="rerun", selection_mode="single-row",
    column_config={
        "grade": GRADE_COLUMN,
        "timestamp": st.column_config.DatetimeColumn("generated", format="DD/MM/YYYY HH:mm"),
    },
)

chosen = cv_table.selection["rows"]

if not chosen:
    st.info("Select a CV above to read it.")
    st.stop()

row = generated.iloc[chosen[0]]
cv = json.loads(row["cv"])

st.divider()
st.subheader(row["title"] or "(untitled)")
st.caption(f"{row['company']} - graded {row['grade']}/100 - "
           f"written in {row['locale']}")

photo = st.radio(
    "Photo", ["With photo", "No photo (ATS-safe)"], horizontal=True,
    help="Applicant tracking systems skip images, and some reject a file that "
         "carries one. Pick the ATS-safe version only when submitting through "
         "an automated system.",
) == "With photo"

blocks = docx_gen.sections(cv)
l10n = docx_gen.load_l10n(row["locale"])

try:
    if photo:
        document = pdf_gen.render_pdf(blocks, l10n)
    else:
        document = docx_gen.render_docx(blocks, l10n, photo=False)
except (FileNotFoundError, RuntimeError) as exc:
    # The template is gitignored, so a fresh clone has only the _example one,
    # and a template can also name a locale ID no values_*.json defines. Both
    # are the user's to fix, and a page that died here would hide the message.
    st.error(f"cannot build the CV: {exc}")
    st.stop()

cols = st.columns(2)

if photo:
    pdf_name = docx_gen.filename(row["company"], row["title"]).replace(
        ".docx", ".pdf")
    cols[0].download_button(
        "Download .pdf", data=document, file_name=pdf_name,
        mime="application/pdf", type="primary",
    )
else:
    cols[0].download_button(
        "Download .docx", data=document,
        file_name=docx_gen.filename(row["company"], row["title"]),
        mime="application/vnd.openxmlformats-officedocument."
             "wordprocessingml.document",
        type="primary",
    )

cols[1].link_button("View job posting", row["url"])

st.divider()

# The preview is read back out of the bytes the button hands over, so it cannot
# disagree with the file -- and it needs to know nothing about placeholders,
# section order or locales, all of which live in the template. Only the docx
# path can be read back this way; the PDF has no preview here.
if photo:
    st.caption("Preview is available for the .docx (ATS-safe) version only.")
else:
    for paragraph in Document(io.BytesIO(document)).paragraphs:
        if paragraph.text.strip():
            st.markdown("".join(f"**{run.text}**" if run.bold else run.text
                                for run in paragraph.runs))

with st.expander("Raw JSON"):
    st.json(cv)
